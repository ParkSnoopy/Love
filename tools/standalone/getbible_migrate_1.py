#!/usr/bin/env python3
"""getbible_migrate_1.py - Convert getBible JSON to XLSX

3 sheets matching format_example/kor_korkrv.xlsx:
  - books:   book_id, osis, eng_name, name, testament, chapters
  - verses:  book_id, chapter, verse, text
  - version: slug, label

Output uses 3-letter ISO 639-3 lang prefix (e.g. eng_kjv.xlsx).
Converts getBible's 2-letter lang codes to 3-letter automatically.

Usage:
    python getbible_migrate_1.py input_dir/ [--output-dir out/]
    python getbible_migrate_1.py download/eng_kjv.json
    python getbible_migrate_1.py download/ --output-dir ./xlsx
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path

BOOKS = [
    (1, "Gen", "Genesis", "OT", 50),
    (2, "Exod", "Exodus", "OT", 40),
    (3, "Lev", "Leviticus", "OT", 27),
    (4, "Num", "Numbers", "OT", 36),
    (5, "Deut", "Deuteronomy", "OT", 34),
    (6, "Josh", "Joshua", "OT", 24),
    (7, "Judg", "Judges", "OT", 21),
    (8, "Ruth", "Ruth", "OT", 4),
    (9, "1Sam", "1 Samuel", "OT", 31),
    (10, "2Sam", "2 Samuel", "OT", 24),
    (11, "1Kgs", "1 Kings", "OT", 22),
    (12, "2Kgs", "2 Kings", "OT", 25),
    (13, "1Chr", "1 Chronicles", "OT", 29),
    (14, "2Chr", "2 Chronicles", "OT", 36),
    (15, "Ezra", "Ezra", "OT", 10),
    (16, "Neh", "Nehemiah", "OT", 13),
    (17, "Esth", "Esther", "OT", 10),
    (18, "Job", "Job", "OT", 42),
    (19, "Ps", "Psalms", "OT", 150),
    (20, "Prov", "Proverbs", "OT", 31),
    (21, "Eccl", "Ecclesiastes", "OT", 12),
    (22, "Song", "Song of Songs", "OT", 8),
    (23, "Isa", "Isaiah", "OT", 66),
    (24, "Jer", "Jeremiah", "OT", 52),
    (25, "Lam", "Lamentations", "OT", 5),
    (26, "Ezek", "Ezekiel", "OT", 48),
    (27, "Dan", "Daniel", "OT", 12),
    (28, "Hos", "Hosea", "OT", 14),
    (29, "Joel", "Joel", "OT", 3),
    (30, "Amos", "Amos", "OT", 9),
    (31, "Obad", "Obadiah", "OT", 1),
    (32, "Jonah", "Jonah", "OT", 4),
    (33, "Mic", "Micah", "OT", 7),
    (34, "Nah", "Nahum", "OT", 3),
    (35, "Hab", "Habakkuk", "OT", 3),
    (36, "Zeph", "Zephaniah", "OT", 3),
    (37, "Hag", "Haggai", "OT", 2),
    (38, "Zech", "Zechariah", "OT", 14),
    (39, "Mal", "Malachi", "OT", 4),
    (40, "Matt", "Matthew", "NT", 28),
    (41, "Mark", "Mark", "NT", 16),
    (42, "Luke", "Luke", "NT", 24),
    (43, "John", "John", "NT", 21),
    (44, "Acts", "Acts", "NT", 28),
    (45, "Rom", "Romans", "NT", 16),
    (46, "1Cor", "1 Corinthians", "NT", 16),
    (47, "2Cor", "2 Corinthians", "NT", 13),
    (48, "Gal", "Galatians", "NT", 6),
    (49, "Eph", "Ephesians", "NT", 6),
    (50, "Phil", "Philippians", "NT", 4),
    (51, "Col", "Colossians", "NT", 4),
    (52, "1Thess", "1 Thessalonians", "NT", 5),
    (53, "2Thess", "2 Thessalonians", "NT", 3),
    (54, "1Tim", "1 Timothy", "NT", 6),
    (55, "2Tim", "2 Timothy", "NT", 4),
    (56, "Titus", "Titus", "NT", 3),
    (57, "Phlm", "Philemon", "NT", 1),
    (58, "Heb", "Hebrews", "NT", 13),
    (59, "Jas", "James", "NT", 5),
    (60, "1Pet", "1 Peter", "NT", 5),
    (61, "2Pet", "2 Peter", "NT", 3),
    (62, "1John", "1 John", "NT", 5),
    (63, "2John", "2 John", "NT", 1),
    (64, "3John", "3 John", "NT", 1),
    (65, "Jude", "Jude", "NT", 1),
    (66, "Rev", "Revelation", "NT", 22),
]

BOOK_MAP = {b[0]: b for b in BOOKS}

# ISO 639-1 (2-letter) → ISO 639-3 (3-letter) — matches helloao convention
ISO_639_1_TO_3 = {
    "aa": "aar", "ab": "abk", "ae": "ave", "af": "afr", "ak": "aka",
    "am": "amh", "an": "arg", "ar": "ara", "as": "asm", "av": "ava",
    "ay": "aym", "az": "aze", "ba": "bak", "be": "bel", "bg": "bul",
    "bh": "bih", "bi": "bis", "bm": "bam", "bn": "ben", "bo": "bod",
    "br": "bre", "bs": "bos", "ca": "cat", "ce": "che", "ch": "cha",
    "co": "cos", "cr": "cre", "cs": "ces", "cu": "chu", "cv": "chv",
    "cy": "cym", "da": "dan", "de": "deu", "dv": "div", "dz": "dzo",
    "ee": "ewe", "el": "ell", "en": "eng", "eo": "epo", "es": "spa",
    "et": "est", "eu": "eus", "fa": "fas", "ff": "ful", "fi": "fin",
    "fj": "fij", "fo": "fao", "fr": "fra", "fy": "fry", "ga": "gle",
    "gd": "gla", "gl": "glg", "gn": "grn", "gu": "guj", "gv": "glv",
    "ha": "hau", "he": "heb", "hi": "hin", "ho": "hmo", "hr": "hrv",
    "ht": "hat", "hu": "hun", "hy": "hye", "hz": "her", "ia": "ina",
    "id": "ind", "ie": "ile", "ig": "ibo", "ii": "iii", "ik": "ipk",
    "io": "ido", "is": "isl", "it": "ita", "iu": "iku", "ja": "jpn",
    "jv": "jav", "ka": "kat", "kg": "kon", "ki": "kik", "kj": "kua",
    "kk": "kaz", "kl": "kal", "km": "khm", "kn": "kan", "ko": "kor",
    "kr": "kau", "ks": "kas", "ku": "kur", "kv": "kom", "kw": "cor",
    "ky": "kir", "la": "lat", "lb": "ltz", "lg": "lug", "li": "lim",
    "ln": "lin", "lo": "lao", "lt": "lit", "lu": "lub", "lv": "lav",
    "mg": "mlg", "mh": "mah", "mi": "mri", "mk": "mkd", "ml": "mal",
    "mn": "mon", "mr": "mar", "ms": "msa", "mt": "mlt", "my": "mya",
    "na": "nau", "nb": "nob", "nd": "nde", "ne": "nep", "ng": "ndo",
    "nl": "nld", "nn": "nno", "no": "nor", "nr": "nbl", "nv": "nav",
    "ny": "nya", "oc": "oci", "oj": "oji", "om": "orm", "or": "ori",
    "os": "oss", "pa": "pan", "pi": "pli", "pl": "pol", "ps": "pus",
    "pt": "por", "qu": "que", "rm": "roh", "rn": "run", "ro": "ron",
    "ru": "rus", "rw": "kin", "sa": "san", "sc": "srd", "sd": "snd",
    "se": "sme", "sg": "sag", "si": "sin", "sk": "slk", "sl": "slv",
    "sm": "smo", "sn": "sna", "so": "som", "sq": "sqi", "sr": "srp",
    "ss": "ssw", "st": "sot", "su": "sun", "sv": "swe", "sw": "swa",
    "ta": "tam", "te": "tel", "tg": "tgk", "th": "tha", "ti": "tir",
    "tk": "tuk", "tl": "tgl", "tn": "tsn", "to": "ton", "tr": "tur",
    "ts": "tso", "tt": "tat", "tw": "twi", "ty": "tah", "ug": "uig",
    "uk": "ukr", "ur": "urd", "uz": "uzb", "ve": "ven", "vi": "vie",
    "vo": "vol", "wa": "wln", "wo": "wol", "xh": "xho", "yi": "yid",
    "yo": "yor", "za": "zha", "zh": "zho", "zu": "zul",
}


def _resolve_label(data, use_localized=False):
    """Get version label from getBible JSON fields.
    Default: use translation (compact English name).
    --localized: use description (often native/localized, may be verbose).
    """
    if use_localized:
        return data.get("description") or data.get("translation", "")
    return data.get("translation") or data.get("description", "")


def convert_json_to_xlsx(json_path, output_dir, use_localized=False):
    import openpyxl

    try:
        with open(json_path, "rb") as f:
            data = json.load(f)
    except (json.JSONDecodeError, OSError) as e:
        print(f"  [SKIP] JSON error: {json_path}: {e}", flush=True)
        return False

    if not isinstance(data, dict) or "books" not in data:
        print(f"  [SKIP] Not a getBible full translation JSON: {json_path}", flush=True)
        return False

    abbr = data.get("abbreviation") or json_path.stem
    lang2 = data.get("lang", "unknown").split("-")[0].lower()  # strip region (zh-Hans → zh)
    lang3 = ISO_639_1_TO_3.get(lang2, lang2)  # convert to 3-letter
    slug = abbr
    label = _resolve_label(data, use_localized)

    out_name = f"{lang3}_{slug}.xlsx"
    out_path = output_dir / out_name

    if out_path.exists():
        print(f"  EXISTS: {out_name}", flush=True)
        return True

    books_list = data.get("books", [])
    if isinstance(books_list, dict):
        books_list = list(books_list.values())

    if not books_list:
        print(f"  [SKIP] No books data: {json_path}", flush=True)
        return False

    # Parse book names from JSON
    book_names = {}
    verses_rows = []

    for bk in books_list:
        if isinstance(bk, dict):
            nr = bk.get("nr") or bk.get("book_nr")
            if nr is None:
                continue
            nr = int(nr)
            bname = bk.get("name", "")
            if bname:
                book_names[nr] = bname

            chapters = bk.get("chapters", [])
            if isinstance(chapters, dict):
                chapters = list(chapters.values())

            for ch_data in chapters:
                if isinstance(ch_data, dict):
                    ch_num = ch_data.get("chapter", 0)
                    verses = ch_data.get("verses", [])
                    if isinstance(verses, dict):
                        verses = list(verses.values())
                    for v in verses:
                        if isinstance(v, dict):
                            vnum = v.get("verse", 0)
                            text = (v.get("text", "") or "").strip()
                            if text:
                                verses_rows.append((nr, ch_num, vnum, text))
                        elif isinstance(v, str):
                            # chapters might be list of strings
                            vnum = verses.index(v) + 1
                            if v.strip():
                                verses_rows.append((nr, ch_num, vnum, v.strip()))

    if not verses_rows:
        print(f"  [SKIP] No verses found: {json_path}", flush=True)
        return False

    output_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    # Sheet 1: books
    ws_books = wb.active
    ws_books.title = "books"
    ws_books.append(("book_id", "osis", "eng_name", "name", "testament", "chapters"))
    for bid, osis, eng_name, testament, chapters in BOOKS:
        localized_name = book_names.get(bid, "")
        ws_books.append((bid, osis, eng_name, localized_name, testament, chapters))

    # Sheet 2: verses
    ws_verses = wb.create_sheet("verses")
    ws_verses.append(("book_id", "chapter", "verse", "text"))
    verses_rows.sort(key=lambda r: (r[0], r[1], r[2]))
    for row in verses_rows:
        ws_verses.append(row)

    # Sheet 3: version
    ws_ver = wb.create_sheet("version")
    ws_ver.append(("slug", "label"))
    ws_ver.append((slug, label))

    wb.save(out_path)
    print(
        f"  OK: {out_path}  ({len(verses_rows)} verses, {len(book_names)} books)",
        flush=True,
    )
    return True


def collect_json_files(path):
    if path.is_file():
        if path.suffix.lower() == ".json":
            return [path]
        print(f"  [SKIP] Not JSON: {path}", flush=True)
        return []
    if path.is_dir():
        return sorted(
            p for p in path.iterdir() if p.is_file() and p.suffix.lower() == ".json"
        )
    print(f"  [SKIP] Not found: {path}", flush=True)
    return []


def main():
    ap = argparse.ArgumentParser(
        description="Convert getBible JSON file(s) to XLSX.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    ap.add_argument(
        "input",
        nargs="+",
        help="getBible JSON file(s) or directories containing .json files",
    )
    ap.add_argument(
        "--output-dir",
        "-o",
        default="./xlsx",
        help="Output directory (default: ./xlsx)",
    )
    ap.add_argument(
        "--localized",
        action="store_true",
        help="Use description field (localized name, may be verbose) instead of translation",
    )
    args = ap.parse_args()

    output_dir = Path(args.output_dir)

    json_files = []
    for inp in args.input:
        json_files.extend(collect_json_files(Path(inp)))
    json_files = sorted(set(json_files))
    if not json_files:
        print("No JSON files found.", flush=True)
        sys.exit(1)

    ok = 0
    for jsf in json_files:
        if convert_json_to_xlsx(jsf, output_dir, args.localized):
            ok += 1

    print(f"\nDone: {ok}/{len(json_files)} converted -> {output_dir}/", flush=True)
    sys.exit(0 if ok == len(json_files) else 1)


if __name__ == "__main__":
    main()
