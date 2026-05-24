#!/usr/bin/env python3
"""helloao_migrate_1.py - Convert helloao complete.json to XLSX

3 sheets matching format_example/kor_korkrv.xlsx:
  - books:   book_id, osis, eng_name, name, testament, chapters
  - verses:  book_id, chapter, verse, text
  - version: slug, label

Naming: <lang>_<id>.xlsx

Usage:
    python helloao_migrate_1.py helloao/BSB.json
    python helloao_migrate_1.py helloao/ --output-dir ./xlsx
"""

import argparse
import json
import os
import sys
from pathlib import Path

BOOKS = [
    (1, "Gen", "Genesis", "OT", 50),
    (2, "Exod", "Exodus", "OT", 40),
    (3, "Lev", "Leviticus", "OT", 27),
    (4, "Num", "Numbers", "OT", 36),
    (5, "Deut", "Deuteronomy", "OT", 34),
    (6, "Josh", "Joshua", "OT", 24),
    (7, "Judg", "Judges", "OT", 21),
    (8, "Ruth", "Ruth", "OT", 4),
    (9, "1Sam", "1 Samuel", "OT", 31),
    (10, "2Sam", "2 Samuel", "OT", 24),
    (11, "1Kgs", "1 Kings", "OT", 22),
    (12, "2Kgs", "2 Kings", "OT", 25),
    (13, "1Chr", "1 Chronicles", "OT", 29),
    (14, "2Chr", "2 Chronicles", "OT", 36),
    (15, "Ezra", "Ezra", "OT", 10),
    (16, "Neh", "Nehemiah", "OT", 13),
    (17, "Esth", "Esther", "OT", 10),
    (18, "Job", "Job", "OT", 42),
    (19, "Ps", "Psalms", "OT", 150),
    (20, "Prov", "Proverbs", "OT", 31),
    (21, "Eccl", "Ecclesiastes", "OT", 12),
    (22, "Song", "Song of Songs", "OT", 8),
    (23, "Isa", "Isaiah", "OT", 66),
    (24, "Jer", "Jeremiah", "OT", 52),
    (25, "Lam", "Lamentations", "OT", 5),
    (26, "Ezek", "Ezekiel", "OT", 48),
    (27, "Dan", "Daniel", "OT", 12),
    (28, "Hos", "Hosea", "OT", 14),
    (29, "Joel", "Joel", "OT", 3),
    (30, "Amos", "Amos", "OT", 9),
    (31, "Obad", "Obadiah", "OT", 1),
    (32, "Jonah", "Jonah", "OT", 4),
    (33, "Mic", "Micah", "OT", 7),
    (34, "Nah", "Nahum", "OT", 3),
    (35, "Hab", "Habakkuk", "OT", 3),
    (36, "Zeph", "Zephaniah", "OT", 3),
    (37, "Hag", "Haggai", "OT", 2),
    (38, "Zech", "Zechariah", "OT", 14),
    (39, "Mal", "Malachi", "OT", 4),
    (40, "Matt", "Matthew", "NT", 28),
    (41, "Mark", "Mark", "NT", 16),
    (42, "Luke", "Luke", "NT", 24),
    (43, "John", "John", "NT", 21),
    (44, "Acts", "Acts", "NT", 28),
    (45, "Rom", "Romans", "NT", 16),
    (46, "1Cor", "1 Corinthians", "NT", 16),
    (47, "2Cor", "2 Corinthians", "NT", 13),
    (48, "Gal", "Galatians", "NT", 6),
    (49, "Eph", "Ephesians", "NT", 6),
    (50, "Phil", "Philippians", "NT", 4),
    (51, "Col", "Colossians", "NT", 4),
    (52, "1Thess", "1 Thessalonians", "NT", 5),
    (53, "2Thess", "2 Thessalonians", "NT", 3),
    (54, "1Tim", "1 Timothy", "NT", 6),
    (55, "2Tim", "2 Timothy", "NT", 4),
    (56, "Titus", "Titus", "NT", 3),
    (57, "Phlm", "Philemon", "NT", 1),
    (58, "Heb", "Hebrews", "NT", 13),
    (59, "Jas", "James", "NT", 5),
    (60, "1Pet", "1 Peter", "NT", 5),
    (61, "2Pet", "2 Peter", "NT", 3),
    (62, "1John", "1 John", "NT", 5),
    (63, "2John", "2 John", "NT", 1),
    (64, "3John", "3 John", "NT", 1),
    (65, "Jude", "Jude", "NT", 1),
    (66, "Rev", "Revelation", "NT", 22),
]

BOOK_MAP = {b[0]: b for b in BOOKS}

# Map helloao book IDs (USFM) to canonical book numbers (1-66)
# Based on book.order from the API (standard Bible ordering)
USFM_TO_BOOK_ID = {
    "GEN": 1,
    "EXO": 2,
    "LEV": 3,
    "NUM": 4,
    "DEU": 5,
    "JOS": 6,
    "JDG": 7,
    "RUT": 8,
    "1SA": 9,
    "2SA": 10,
    "1KI": 11,
    "2KI": 12,
    "1CH": 13,
    "2CH": 14,
    "EZR": 15,
    "NEH": 16,
    "EST": 17,
    "JOB": 18,
    "PSA": 19,
    "PRO": 20,
    "ECC": 21,
    "SNG": 22,
    "ISA": 23,
    "JER": 24,
    "LAM": 25,
    "EZK": 26,
    "DAN": 27,
    "HOS": 28,
    "JOL": 29,
    "AMO": 30,
    "OBA": 31,
    "JON": 32,
    "MIC": 33,
    "NAM": 34,
    "HAB": 35,
    "ZEP": 36,
    "HAG": 37,
    "ZEC": 38,
    "MAL": 39,
    "MAT": 40,
    "MRK": 41,
    "LUK": 42,
    "JHN": 43,
    "ACT": 44,
    "ROM": 45,
    "1CO": 46,
    "2CO": 47,
    "GAL": 48,
    "EPH": 49,
    "PHP": 50,
    "COL": 51,
    "1TH": 52,
    "2TH": 53,
    "1TI": 54,
    "2TI": 55,
    "TIT": 56,
    "PHM": 57,
    "HEB": 58,
    "JAS": 59,
    "1PE": 60,
    "2PE": 61,
    "1JN": 62,
    "2JN": 63,
    "3JN": 64,
    "JUD": 65,
    "REV": 66,
}


def extract_verse_text(content_items):
    """Extract plain text from helloao verse content array.

    Content can contain strings, footnote refs ({noteId: N}),
    line breaks ({lineBreak: true}), inline headings ({heading: "..."}),
    and formatted text ({text: "...", ...}).
    """
    parts = []
    for item in content_items:
        if isinstance(item, str):
            parts.append(item)
        elif isinstance(item, dict):
            if "text" in item:
                parts.append(item["text"])
            elif "heading" in item:
                parts.append(item["heading"])
    return "".join(parts).strip()


def convert_json_to_xlsx(json_path: Path, output_dir: Path) -> bool:
    import openpyxl

    try:
        with open(json_path, "rb") as f:
            data = json.load(f)
    except (json.JSONDecodeError, OSError) as e:
        print(f"  [SKIP] JSON error: {json_path}: {e}", flush=True)
        return False

    if not isinstance(data, dict) or "books" not in data:
        print(
            f"  [SKIP] Not a helloao complete translation JSON: {json_path}", flush=True
        )
        return False

    trans = data.get("translation", {})
    tid = trans.get("id") or json_path.stem
    lang = trans.get("language", "unknown")
    name = trans.get("name", None) or trans.get("englishName", "")

    out_name = f"{lang}_{tid}.xlsx"
    out_path = output_dir / out_name

    if out_path.exists():
        print(f"  EXISTS: {out_name}", flush=True)
        return True

    books_data = data.get("books", [])
    if not books_data:
        print(f"  [SKIP] No books data: {json_path}", flush=True)
        return False

    # Parse books and verses
    book_names = {}
    verses_rows = []

    for bk in books_data:
        bid = bk.get("id", "")
        book_num = USFM_TO_BOOK_ID.get(bid)
        if book_num is None:
            continue

        bname = bk.get("name", "") or bk.get("commonName", "")
        if bname:
            book_names[book_num] = bname

        chapters = bk.get("chapters", [])
        for ch_entry in chapters:
            ch = ch_entry.get("chapter", ch_entry)
            ch_num = ch.get("number", 0)
            content = ch.get("content", [])
            for item in content:
                if not isinstance(item, dict):
                    continue
                if item.get("type") != "verse":
                    continue
                vnum = item.get("number", 0)
                text = extract_verse_text(item.get("content", []))
                if text:
                    verses_rows.append((book_num, ch_num, vnum, text))

    if not verses_rows:
        print(f"  [SKIP] No verses found: {json_path}", flush=True)
        return False

    output_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    # Sheet 1: books
    ws_books = wb.active
    ws_books.title = "books"
    ws_books.append(("book_id", "osis", "eng_name", "name", "testament", "chapters"))
    for bid, osis, eng_name, testament, chapters in BOOKS:
        localized_name = book_names.get(bid, "")
        ws_books.append((bid, osis, eng_name, localized_name, testament, chapters))

    # Sheet 2: verses
    ws_verses = wb.create_sheet("verses")
    ws_verses.append(("book_id", "chapter", "verse", "text"))
    verses_rows.sort(key=lambda r: (r[0], r[1], r[2]))
    for row in verses_rows:
        ws_verses.append(row)

    # Sheet 3: version
    ws_ver = wb.create_sheet("version")
    ws_ver.append(("slug", "label"))
    ws_ver.append((tid, name))

    wb.save(out_path)
    print(
        f"  OK: {out_path}  ({len(verses_rows)} verses, {len(book_names)} books)",
        flush=True,
    )
    return True


def collect_json_files(path: Path) -> list[Path]:
    if path.is_file():
        if path.suffix.lower() == ".json":
            return [path]
        print(f"  [SKIP] Not JSON: {path}", flush=True)
        return []
    if path.is_dir():
        return sorted(
            p for p in path.iterdir() if p.is_file() and p.suffix.lower() == ".json"
        )
    print(f"  [SKIP] Not found: {path}", flush=True)
    return []


def main():
    ap = argparse.ArgumentParser(
        description="Convert helloao complete.json file(s) to XLSX.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    ap.add_argument(
        "input",
        nargs="+",
        help="helloao JSON file(s) or directories containing .json files",
    )
    ap.add_argument(
        "--output-dir",
        "-o",
        default="./xlsx",
        help="Output directory (default: ./xlsx)",
    )
    args = ap.parse_args()

    output_dir = Path(args.output_dir)

    json_files = []
    for inp in args.input:
        json_files.extend(collect_json_files(Path(inp)))
    json_files = sorted(set(json_files))
    if not json_files:
        print("No JSON files found.", flush=True)
        sys.exit(1)

    ok = 0
    for jsf in json_files:
        if convert_json_to_xlsx(jsf, output_dir):
            ok += 1

    print(f"\nDone: {ok}/{len(json_files)} converted -> {output_dir}/", flush=True)
    sys.exit(0 if ok == len(json_files) else 1)


if __name__ == "__main__":
    main()
