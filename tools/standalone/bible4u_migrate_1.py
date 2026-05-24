#!/usr/bin/env python3
"""bible4u_migrate_1.py - Convert Bible4U XML to XLSX

Format: 3 sheets matching format_example/kor_korkrv.xlsx
  - books:   book_id, osis, eng_name, name, testament, chapters
  - verses:  book_id, chapter, verse, text
  - version: slug, label

Naming: <ISO country code>_<bible version name>.xlsx

Usage:
    python bible4u_migrate_1.py input_dir/ [--output-dir out/]
    python bible4u_migrate_1.py download/kor.xml
    python bible4u_migrate_1.py download/ --output-dir ./xlsx
"""

import argparse
import os
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

# ---------------------------------------------------------------------------
# Canonical book list (from extractors.py)
# ---------------------------------------------------------------------------
BOOKS = [
    (1, "Gen", "Genesis", "OT", 50),
    (2, "Exod", "Exodus", "OT", 40),
    (3, "Lev", "Leviticus", "OT", 27),
    (4, "Num", "Numbers", "OT", 36),
    (5, "Deut", "Deuteronomy", "OT", 34),
    (6, "Josh", "Joshua", "OT", 24),
    (7, "Judg", "Judges", "OT", 21),
    (8, "Ruth", "Ruth", "OT", 4),
    (9, "1Sam", "1 Samuel", "OT", 31),
    (10, "2Sam", "2 Samuel", "OT", 24),
    (11, "1Kgs", "1 Kings", "OT", 22),
    (12, "2Kgs", "2 Kings", "OT", 25),
    (13, "1Chr", "1 Chronicles", "OT", 29),
    (14, "2Chr", "2 Chronicles", "OT", 36),
    (15, "Ezra", "Ezra", "OT", 10),
    (16, "Neh", "Nehemiah", "OT", 13),
    (17, "Esth", "Esther", "OT", 10),
    (18, "Job", "Job", "OT", 42),
    (19, "Ps", "Psalms", "OT", 150),
    (20, "Prov", "Proverbs", "OT", 31),
    (21, "Eccl", "Ecclesiastes", "OT", 12),
    (22, "Song", "Song of Songs", "OT", 8),
    (23, "Isa", "Isaiah", "OT", 66),
    (24, "Jer", "Jeremiah", "OT", 52),
    (25, "Lam", "Lamentations", "OT", 5),
    (26, "Ezek", "Ezekiel", "OT", 48),
    (27, "Dan", "Daniel", "OT", 12),
    (28, "Hos", "Hosea", "OT", 14),
    (29, "Joel", "Joel", "OT", 3),
    (30, "Amos", "Amos", "OT", 9),
    (31, "Obad", "Obadiah", "OT", 1),
    (32, "Jonah", "Jonah", "OT", 4),
    (33, "Mic", "Micah", "OT", 7),
    (34, "Nah", "Nahum", "OT", 3),
    (35, "Hab", "Habakkuk", "OT", 3),
    (36, "Zeph", "Zephaniah", "OT", 3),
    (37, "Hag", "Haggai", "OT", 2),
    (38, "Zech", "Zechariah", "OT", 14),
    (39, "Mal", "Malachi", "OT", 4),
    (40, "Matt", "Matthew", "NT", 28),
    (41, "Mark", "Mark", "NT", 16),
    (42, "Luke", "Luke", "NT", 24),
    (43, "John", "John", "NT", 21),
    (44, "Acts", "Acts", "NT", 28),
    (45, "Rom", "Romans", "NT", 16),
    (46, "1Cor", "1 Corinthians", "NT", 16),
    (47, "2Cor", "2 Corinthians", "NT", 13),
    (48, "Gal", "Galatians", "NT", 6),
    (49, "Eph", "Ephesians", "NT", 6),
    (50, "Phil", "Philippians", "NT", 4),
    (51, "Col", "Colossians", "NT", 4),
    (52, "1Thess", "1 Thessalonians", "NT", 5),
    (53, "2Thess", "2 Thessalonians", "NT", 3),
    (54, "1Tim", "1 Timothy", "NT", 6),
    (55, "2Tim", "2 Timothy", "NT", 4),
    (56, "Titus", "Titus", "NT", 3),
    (57, "Phlm", "Philemon", "NT", 1),
    (58, "Heb", "Hebrews", "NT", 13),
    (59, "Jas", "James", "NT", 5),
    (60, "1Pet", "1 Peter", "NT", 5),
    (61, "2Pet", "2 Peter", "NT", 3),
    (62, "1John", "1 John", "NT", 5),
    (63, "2John", "2 John", "NT", 1),
    (64, "3John", "3 John", "NT", 1),
    (65, "Jude", "Jude", "NT", 1),
    (66, "Rev", "Revelation", "NT", 22),
]

BOOK_MAP = {b[0]: b for b in BOOKS}  # book_id -> tuple

# Regex to extract ISO code from language tag:  "한국어 (KOR)" -> "KOR"
_LANG_CODE_RE = re.compile(r"\((\w+)\)")


def extract_iso_code(lang_text: str) -> str:
    """Extract ISO country/language code from '<Language> (CODE)' format."""
    m = _LANG_CODE_RE.search(lang_text)
    return m.group(1).lower() if m else "unknown"


def convert_xml_to_xlsx(xml_path: Path, output_dir: Path) -> bool:
    """Convert one Bible4U XML file to XLSX."""
    import openpyxl  # lazy import for CLI speed

    try:
        tree = ET.parse(xml_path)
    except ET.ParseError as e:
        print(f"  [SKIP] Parse error: {xml_path}: {e}", flush=True)
        return False

    root = tree.getroot()
    info = root.find("INFORMATION")
    if info is None:
        print(f"  [SKIP] No INFORMATION block: {xml_path}", flush=True)
        return False

    # ---- Metadata ----
    lang_el = info.find("language")
    ident_el = info.find("identifier")
    lang_text = lang_el.text if lang_el is not None else ""
    identifier = (
        ident_el.text.strip() if ident_el is not None and ident_el.text else "unknown"
    )
    biblename = root.get("biblename", identifier)

    iso_code = extract_iso_code(lang_text)

    # ---- Filename: <ISO>_<version>.xlsx ----
    out_name = f"{iso_code}_{identifier}.xlsx"
    out_path = output_dir / out_name

    if out_path.exists():
        print(f"  EXISTS: {out_name}", flush=True)
        return True

    # ---- Parse books & verses ----
    book_names: dict[int, str] = {}  # book_id -> localized name from XML
    verses_rows = []
    for bible_book in root.findall("BIBLEBOOK"):
        book_id_str = bible_book.get("bnumber", "0")
        book_id = int(book_id_str)
        # Validate against canonical list
        if book_id not in BOOK_MAP:
            continue
        bname = bible_book.get("bname", "")
        if bname:
            book_names[book_id] = bname
        for chapter in bible_book.findall("CHAPTER"):
            ch_str = chapter.get("cnumber", "0")
            chapter_num = int(ch_str)
            for verse in chapter.findall("VERS"):
                v_str = verse.get("vnumber", "0")
                verse_num = int(v_str)
                text = (verse.text or "").strip()
                if text:
                    verses_rows.append((book_id, chapter_num, verse_num, text))

    if not verses_rows:
        print(f"  [SKIP] No verses found: {xml_path}", flush=True)
        return False

    # ---- Build XLSX ----
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    # Sheet 1: books
    ws_books = wb.active
    ws_books.title = "books"
    ws_books.append(("book_id", "osis", "eng_name", "name", "testament", "chapters"))
    for bid, osis, eng_name, testament, chapters in BOOKS:
        localized_name = book_names.get(bid, "")
        ws_books.append((bid, osis, eng_name, localized_name, testament, chapters))

    # Sheet 2: verses
    ws_verses = wb.create_sheet("verses")
    ws_verses.append(("book_id", "chapter", "verse", "text"))
    verses_rows.sort(key=lambda r: (r[0], r[1], r[2]))
    for row in verses_rows:
        ws_verses.append(row)

    # Sheet 3: version
    ws_ver = wb.create_sheet("version")
    ws_ver.append(("slug", "label"))
    ws_ver.append((identifier, biblename))

    wb.save(out_path)
    print(
        f"  OK: {out_path}  ({len(verses_rows)} verses, {len(BOOKS)} books)", flush=True
    )
    return True


def collect_xml_files(path: Path) -> list[Path]:
    """Return all .xml files (single file or directory scan)."""
    if path.is_file():
        if path.suffix.lower() == ".xml":
            return [path]
        print(f"  [SKIP] Not XML: {path}", flush=True)
        return []
    if path.is_dir():
        return sorted(
            p for p in path.iterdir() if p.is_file() and p.suffix.lower() == ".xml"
        )
    print(f"  [SKIP] Not found: {path}", flush=True)
    return []


def main():
    ap = argparse.ArgumentParser(
        description="Convert Bible4U XML file(s) to XLSX.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    ap.add_argument("input", help="Bible4U XML file or directory containing .xml files")
    ap.add_argument(
        "--output-dir",
        "-o",
        default="./xlsx",
        help="Output directory (default: ./xlsx)",
    )
    args = ap.parse_args()

    input_path = Path(args.input)
    output_dir = Path(args.output_dir)

    xml_files = collect_xml_files(input_path)
    if not xml_files:
        print("No XML files found.", flush=True)
        sys.exit(1)

    ok = 0
    for xml_file in xml_files:
        if convert_xml_to_xlsx(xml_file, output_dir):
            ok += 1

    print(f"\nDone: {ok}/{len(xml_files)} converted -> {output_dir}/", flush=True)
    sys.exit(0 if ok == len(xml_files) else 1)


if __name__ == "__main__":
    main()
