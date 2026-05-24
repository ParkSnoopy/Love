#!/usr/bin/env python3
"""getbible_migrate_2.py - Convert getBible XLSX (from getbible_migrate_1) to SQLite

Usage:
    python getbible_migrate_2.py en_kjv.xlsx
    python getbible_migrate_2.py xlsx/ --output-dir ./db
    python getbible_migrate_2.py xlsx/ --single-db bibles.db
"""

import argparse
import os
import sqlite3
import sys
from pathlib import Path

import openpyxl


def parse_slug_from_name(xlsx_path: Path) -> str:
    stem = xlsx_path.stem
    parts = stem.split("_", 1)
    return parts[-1] if len(parts) > 1 else stem


def create_tables(conn):
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS books (
            book_id INTEGER PRIMARY KEY,
            osis TEXT NOT NULL,
            eng_name TEXT NOT NULL,
            name TEXT NOT NULL,
            testament TEXT NOT NULL CHECK(testament IN ('OT', 'NT')),
            chapters INTEGER NOT NULL
        );

        CREATE TABLE IF NOT EXISTS verses (
            book_id INTEGER NOT NULL,
            chapter INTEGER NOT NULL,
            verse INTEGER NOT NULL,
            text TEXT NOT NULL,
            PRIMARY KEY (book_id, chapter, verse),
            FOREIGN KEY (book_id) REFERENCES books(book_id)
        );

        CREATE INDEX IF NOT EXISTS idx_verses_book_chapter
            ON verses(book_id, chapter);

        CREATE TABLE IF NOT EXISTS version (
            slug TEXT PRIMARY KEY,
            label TEXT NOT NULL
        );

        CREATE VIEW IF NOT EXISTS v_verses AS
            SELECT
                v.book_id,
                b.osis,
                b.eng_name,
                b.name AS book_name,
                b.testament,
                v.chapter,
                v.verse,
                v.text
            FROM verses v
            JOIN books b ON b.book_id = v.book_id
            ORDER BY v.book_id, v.chapter, v.verse;
    """)


def import_xlsx(xlsx_path: Path, conn) -> int:
    slug = parse_slug_from_name(xlsx_path)

    wb = openpyxl.load_workbook(xlsx_path, read_only=True)

    ws = wb["books"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn.executemany(
        "INSERT OR REPLACE INTO books (book_id, osis, eng_name, name, testament, chapters) VALUES (?, ?, ?, ?, ?, ?)",
        rows,
    )

    ws3 = wb["version"]
    for row in ws3.iter_rows(min_row=2, values_only=True):
        slug_val, label = row
        conn.execute(
            "INSERT OR REPLACE INTO version (slug, label) VALUES (?, ?)",
            (slug_val or slug, label or slug),
        )

    ws2 = wb["verses"]
    batch = []
    total = 0
    for row in ws2.iter_rows(min_row=2, values_only=True):
        bid, ch, vnum, text = row
        batch.append((bid, ch, vnum, text))
        total += 1
        if len(batch) >= 5000:
            conn.executemany(
                "INSERT OR REPLACE INTO verses (book_id, chapter, verse, text) VALUES (?, ?, ?, ?)",
                batch,
            )
            batch = []
    if batch:
        conn.executemany(
            "INSERT OR REPLACE INTO verses (book_id, chapter, verse, text) VALUES (?, ?, ?, ?)",
            batch,
        )

    wb.close()
    return total


def convert_xlsx_to_sqlite(xlsx_path: Path, output_dir: Path, single_conn=None):
    slug = parse_slug_from_name(xlsx_path)
    db_name = f"{xlsx_path.stem}.sqlite"
    db_path = output_dir / db_name

    if not single_conn and db_path.exists():
        print(f"  EXISTS: {db_name}", flush=True)
        return True

    own_conn = single_conn is None
    conn = single_conn or sqlite3.connect(str(db_path))

    try:
        if own_conn:
            conn.execute("PRAGMA journal_mode=WAL")
            conn.execute("PRAGMA synchronous=OFF")
            conn.execute("PRAGMA cache_size=-80000")

        create_tables(conn)

        total = import_xlsx(xlsx_path, conn)

        if own_conn:
            conn.commit()
            db_size = db_path.stat().st_size
            print(
                f"  OK: {db_name}  ({total} verses, {db_size / 1024 / 1024:.1f} MB)",
                flush=True,
            )
            conn.close()
    except Exception as e:
        if own_conn:
            conn.close()
        raise e

    return True


def collect_xlsx_files(path: Path) -> list[Path]:
    if path.is_file():
        if path.suffix.lower() in (".xlsx", ".xls"):
            return [path]
        print(f"  [SKIP] Not XLSX: {path}", flush=True)
        return []
    if path.is_dir():
        return sorted(
            p for p in path.iterdir() if p.is_file() and p.suffix.lower() == ".xlsx"
        )
    print(f"  [SKIP] Not found: {path}", flush=True)
    return []


def main():
    ap = argparse.ArgumentParser(
        description="Convert getBible XLSX file(s) to SQLite database.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    ap.add_argument(
        "input", nargs="+", help="XLSX file(s) or directories containing .xlsx files"
    )
    ap.add_argument(
        "--output-dir", "-o", default="./db", help="Output directory (default: ./db)"
    )

    ap.add_argument(
        "--single-db",
        metavar="FILE",
        help="Import all XLSX into a single SQLite database instead of one DB per file",
    )

    args = ap.parse_args()
    output_dir = Path(args.output_dir)

    xlsx_files = []
    for inp in args.input:
        xlsx_files.extend(collect_xlsx_files(Path(inp)))
    xlsx_files = sorted(set(xlsx_files))
    if not xlsx_files:
        print("No XLSX files found.", flush=True)
        sys.exit(1)

    print(f"Found {len(xlsx_files)} XLSX file(s).", flush=True)

    single_conn = None
    if args.single_db:
        db_path = Path(args.single_db)
        db_path.parent.mkdir(parents=True, exist_ok=True)
        single_conn = sqlite3.connect(str(db_path))
        single_conn.execute("PRAGMA journal_mode=WAL")
        single_conn.execute("PRAGMA synchronous=OFF")
        single_conn.execute("PRAGMA cache_size=-80000")
        create_tables(single_conn)
        print(f"Single DB: {db_path}", flush=True)
    else:
        output_dir.mkdir(parents=True, exist_ok=True)

    ok = 0
    for f in xlsx_files:
        try:
            convert_xlsx_to_sqlite(f, output_dir, single_conn=single_conn)
            ok += 1
        except Exception as e:
            print(f"  FAIL: {f}: {e}", flush=True)

    if single_conn:
        single_conn.commit()
        db_path = Path(args.single_db)
        db_size = db_path.stat().st_size
        print(
            f"\nDone: {ok}/{len(xlsx_files)} imported -> {db_path} ({db_size / 1024 / 1024:.1f} MB)",
            flush=True,
        )
        single_conn.close()
    else:
        print(f"\nDone: {ok}/{len(xlsx_files)} converted -> {output_dir}/", flush=True)

    sys.exit(0 if ok == len(xlsx_files) else 1)


if __name__ == "__main__":
    main()
